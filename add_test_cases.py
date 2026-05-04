"""Script to replace the Excel test cases with a new dataset."""
import openpyxl
from openpyxl.styles import Alignment
from pathlib import Path

# Test cases data
TEST_CASES = [
    ("Neg_001", "S", "ubaa kohmd bn", "උඹ කොහොමද බන්"),
    ("Neg_002", "S", "mta hrtya thern na", "මට හරියට තේරුණේ නෑ"),
    ("Neg_003", "S", "oya kawd enne", "ඔයා කවද්ද එන්නේ"),
    ("Neg_004", "S", "mata pissu 😂😂", "මට පිස්සු 😂😂"),
    ("Neg_005", "S", "api mtng ekata ymu", "අපි meeting එකට යමු"),
    ("Neg_006", "S", "mama ynwa oyath enwda", "මම යනවා ඔයත් එනවද"),
    ("Neg_007", "S", "ane mage bk eka denna", "අනේ මගේ පොත දෙන්න"),
    ("Neg_008", "S", "hari hari mama enm", "හරි හරි මම එන්නම්"),
    ("Neg_009", "S", "oyge pwd eka mkdda", "ඔයාගේ password එක මොකක්ද"),
    ("Neg_010", "S", "api trip ekta sg ymu", "අපි trip එකට Singapore යමු"),
    ("Neg_011", "M", "mama schl giya psse gedra awa", "මම school ගියා පස්සේ ගෙදර ආවා"),
    ("Neg_012", "M", "api restrnt ekta gihn kanna", "අපි restaurant එකට ගිහින් කන්න"),
    ("Neg_013", "M", "mta num eka denna plwnda", "මට number එක දෙන්න පුළුවන්ද"),
    ("Neg_014", "M", "api mtch eka blnwa tv ekn", "අපි match එක බලනවා TV එකෙන්"),
    ("Neg_015", "M", "mama exam ekk krla awa", "මම exam එකක් කරලා ආවා"),
    ("Neg_016", "M", "oyta meka krnna plwnda", "ඔයාට මේක කරන්න පුළුවන්ද"),
    ("Neg_017", "M", "api dn on d way inne", "අපි දැන් on the way ඉන්නේ"),
    ("Neg_018", "M", "mama msg ekk damma", "මම message එකක් දැම්මා"),
    ("Neg_019", "M", "oyge frnd enwa kiwwa", "ඔයාගේ friend එනවා කිව්වා"),
    ("Neg_020", "M", "mama film ekk blnwa gdr", "මම film එකක් බලනවා"),
    ("Neg_021", "M", "api beach ekta ymu wknd eke", "අපි beach එකට යමු weekend එකේ"),
    ("Neg_022", "M", "mta ktha krnna one", "මට කතා කරන්න ඕන"),
    ("Neg_023", "M", "oyta mkdda hithnne", "ඔයාට මොකක්ද හිතෙන්නේ"),
    ("Neg_024", "M", "mama gym giya psse tired una", "මම gym ගියා පස්සේ tired උනා"),
    ("Neg_025", "M", "api trip plan krmu nxt mnth", "අපි trip plan කරමු next month"),
    ("Neg_026", "M", "laptop eka wda krnwda", "laptop එක වැඩ කරනවද"),
    ("Neg_027", "M", "link ekk ewwa blnna", "link එකක් එව්වා බලන්න"),
    ("Neg_028", "M", "cls late una nsa sir bninna plwn", "class late උනා නිසා sir බැනින්න පුළුවන්"),
    ("Neg_029", "M", "call krnna try kla", "call කරන්න try කළා"),
    ("Neg_030", "M", "meka thrungnna amruda", "මේක තේරුම්ගන්න අමාරුද"),
    ("Neg_031", "L", "mama schl giya psse frnds ekka mtch gahla gedra awa", "school ගියා පස්සේ friends එක්ක match ගහලා ගෙදර ආවා"),
    ("Neg_032", "L", "trip plan krnne sg yanna nxt mnth", "trip plan කරන්නේ Singapore යන්න next month"),
    ("Neg_033", "L", "msg damma hlp one kiyala", "message දැම්මා help ඕන කියලා"),
    ("Neg_034", "L", "restrnt gihn movie blnna plan krla inne", "restaurant ගිහින් movie බලන්න plan කරලා ඉන්නේ"),
    ("Neg_035", "L", "office giya psse mtng tibba busy una", "office ගියා පස්සේ meeting තිබ්බ busy උනා"),
    ("Neg_036", "L", "system use krnna amruda explain krnna plwn", "system use කරන්න අමාරුද explain කරන්න පුළුවන්"),
    ("Neg_037", "L", "beach ynna plan krla oyath enna", "beach යන්න plan කරලා ඔයත් එන්න"),
    ("Neg_038", "L", "call kla nmuth phn off", "call කළා නමුත් phone off"),
    ("Neg_039", "L", "cls late nsa sir bninwa kalin enna one", "class late නිසා sir බනිනවා කලින් එන්න ඕන"),
    ("Neg_040", "L", "support one kiyala msg damma reply denna", "support ඕන කියලා message දැම්මා reply දෙන්න"),
    ("Neg_041", "S", "ub kohmda 😅", "උඹ කොහොමද 😅"),
    ("Neg_042", "S", "mta enna ba bn", "මට එන්න බෑ බන්"),
    ("Neg_043", "S", "oyla kohmda inne", "ඔයාලා කොහොමද ඉන්නේ"),
    ("Neg_044", "S", "mama lunch gtta 🍔", "මම lunch ගත්තා 🍔"),
    ("Neg_045", "S", "oyta hithnwda", "ඔයාට හිතෙනවද"),
    ("Neg_046", "M", "mama wifi thynwda baluwa", "මම WiFi තියෙනවද බලුවා"),
    ("Neg_047", "M", "api bus ekn ymu", "අපි bus එකෙන් යමු"),
    ("Neg_048", "M", "oyge email eka denna plz", "ඔයාගේ email එක දෙන්න"),
    ("Neg_049", "M", "mama game eka gahla inne 🎮", "මම game එක ගහලා ඉන්නේ 🎮"),
    ("Neg_050", "M", "oyta call ekk denna puluwanda", "ඔයාට call එකක් දෙන්න පුළුවන්ද"),
]

def add_test_cases_to_excel():
    """Replace the workbook rows with the configured test cases."""
    excel_path = Path(__file__).parent / "Assignment 1 - Test cases.xlsx"
    
    if not excel_path.exists():
        print(f"Error: Excel file not found at {excel_path}")
        return False
    
    try:
        wb = openpyxl.load_workbook(excel_path)
        ws = wb.active
        
        print(f"Loaded worksheet: {ws.title}")
        print(f"Current max row: {ws.max_row}")

        header_row = 1
        if ws.max_row > 0:
            first_cell = ws.cell(1, 1).value
            if first_cell and "TC" in str(first_cell).upper():
                header_row = 1
            else:
                for row in range(1, min(10, ws.max_row + 1)):
                    cell_val = ws.cell(row, 1).value
                    if cell_val and "TC" in str(cell_val).upper():
                        header_row = row
                        break

        if ws.max_row > header_row:
            ws.delete_rows(header_row + 1, ws.max_row - header_row)

        start_row = header_row + 1
        
        print(f"Header row: {header_row}")
        print(f"Replacing test cases starting at row: {start_row}")
        
        for idx, (tc_id, length_type, input_text, expected_output) in enumerate(TEST_CASES, start=1):
            row = start_row + idx - 1
            ws.cell(row=row, column=1).value = tc_id
            ws.cell(row=row, column=2).value = length_type
            ws.cell(row=row, column=3).value = input_text
            ws.cell(row=row, column=4).value = expected_output
            ws.cell(row=row, column=5).value = None
            ws.cell(row=row, column=6).value = None
            
            for col in range(1, 7):
                cell = ws.cell(row=row, column=col)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        
        wb.save(excel_path)
        print(f"✓ Successfully replaced workbook data with {len(TEST_CASES)} test cases in {excel_path}")
        print(f"Test cases written in rows {start_row} to {start_row + len(TEST_CASES) - 1}")
        return True
        
    except Exception as e:
        print(f"Error: {e}")
        return False

if __name__ == "__main__":
    success = add_test_cases_to_excel()
    exit(0 if success else 1)
